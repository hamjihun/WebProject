#!/usr/bin/env python3
"""
PPTX 차트의 외부 엑셀 링크를 내장 데이터로 변환한다.

외부 파일(예: J:\\...\\손익 엑셀 자료.xlsx)에 연결된 차트는 그 파일이 없는 PC에서
"데이터 편집"이 되지 않는다. 이 스크립트는 차트 XML 안에 캐시된 값(현재 그래프에
표시되는 값)을 그대로 읽어 차트마다 작은 워크북(.xlsx)을 만들어 PPTX 안에 내장하고,
링크를 그 워크북으로 바꾼다. 셀 주소·시트 이름은 원본 수식(c:f)과 동일하게 유지하므로
차트 서식과 모양은 전혀 바뀌지 않는다.

사용법:
    python3 tools/embed_chart_data.py 입력.pptx 출력.pptx

필요 패키지: lxml, openpyxl
"""
import os
import re
import shutil
import sys
import tempfile
import zipfile

from lxml import etree
from openpyxl import Workbook
from openpyxl.utils import range_boundaries

NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
REL_PACKAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
REL_OLE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def split_ref(formula):
    """'시트 이름'!$A$1:$B$2 -> ('시트 이름', 'A1:B2')"""
    m = re.match(r"^(?:'((?:[^']|'')*)'|([^'!]+))!(.+)$", formula)
    if not m:
        raise ValueError(f"지원하지 않는 참조: {formula}")
    sheet = (m.group(1) or m.group(2)).replace("''", "'")
    return sheet, m.group(3).replace("$", "")


def cells_in(rng):
    """차트 포인트 순서(idx)와 같은 순서로 셀 좌표를 돌려준다."""
    c1, r1, c2, r2 = range_boundaries(rng if ":" in rng else f"{rng}:{rng}")
    if r1 == r2:
        return [(r1, c) for c in range(c1, c2 + 1)]
    if c1 == c2:
        return [(r, c1) for r in range(r1, r2 + 1)]
    return [(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def to_number(text):
    try:
        v = float(text)
    except ValueError:
        return text
    if v.is_integer() and "." not in text and "E" not in text.upper():
        return int(v)
    return v


def build_workbook(chart_root):
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {}
    written = 0
    for ref in chart_root.iter("{%s}strRef" % NS["c"], "{%s}numRef" % NS["c"]):
        f_el = ref.find("c:f", NS)
        if f_el is None or not f_el.text:
            continue
        sheet, rng = split_ref(f_el.text)
        ws = sheets.get(sheet)
        if ws is None:
            ws = sheets[sheet] = wb.create_sheet(title=sheet)
        is_num = ref.tag.endswith("numRef")
        cache = ref.find("c:numCache" if is_num else "c:strCache", NS)
        if cache is None:
            continue
        fmt_el = cache.find("c:formatCode", NS)
        fmt = fmt_el.text if fmt_el is not None else None
        pts = {int(p.get("idx")): p.findtext("c:v", namespaces=NS) for p in cache.findall("c:pt", NS)}
        for i, (row, col) in enumerate(cells_in(rng)):
            cell = ws.cell(row=row, column=col)
            if fmt:
                cell.number_format = fmt
            if i in pts and pts[i] is not None:
                cell.value = to_number(pts[i]) if is_num else pts[i]
                written += 1
    return wb, written


def convert(src, dst):
    work = tempfile.mkdtemp(prefix="pptx_embed_")
    try:
        with zipfile.ZipFile(src) as z:
            names = [i.filename for i in z.infolist()]
            z.extractall(work)
        chart_dir = os.path.join(work, "ppt", "charts")
        emb_dir = os.path.join(work, "ppt", "embeddings")
        os.makedirs(emb_dir, exist_ok=True)
        added = []
        chart_files = sorted(
            (f for f in os.listdir(chart_dir) if re.fullmatch(r"chart\d+\.xml", f)),
            key=lambda f: int(re.findall(r"\d+", f)[0]),
        )
        for cf in chart_files:
            n = int(re.findall(r"\d+", cf)[0])
            cpath = os.path.join(chart_dir, cf)
            rpath = os.path.join(chart_dir, "_rels", cf + ".rels")
            tree = etree.parse(cpath)
            root = tree.getroot()
            ext = root.find("c:externalData", NS)
            if ext is None or not os.path.exists(rpath):
                print(f"{cf}: 외부 데이터 없음, 건너뜀")
                continue
            rid = ext.get("{%s}id" % NS["r"])
            rtree = etree.parse(rpath)
            rel = next((e for e in rtree.getroot() if e.get("Id") == rid), None)
            if rel is None or rel.get("TargetMode") != "External":
                print(f"{cf}: 이미 내장 데이터, 건너뜀")
                continue
            wb, written = build_workbook(root)
            if written == 0:
                print(f"{cf}: 캐시 데이터 없음, 건너뜀")
                continue
            xlsx_name = f"Microsoft_Excel_Worksheet{n}.xlsx"
            wb.save(os.path.join(emb_dir, xlsx_name))
            rel.set("Type", REL_PACKAGE)
            rel.set("Target", f"../embeddings/{xlsx_name}")
            del rel.attrib["TargetMode"]
            rtree.write(rpath, xml_declaration=True, encoding="UTF-8", standalone=True)
            auto = ext.find("c:autoUpdate", NS)
            if auto is not None:
                auto.set("val", "0")
            tree.write(cpath, xml_declaration=True, encoding="UTF-8", standalone=True)
            added.append(f"ppt/embeddings/{xlsx_name}")
            print(f"{cf}: {wb.sheetnames} 셀 {written}개 내장")

        ct_path = os.path.join(work, "[Content_Types].xml")
        with open(ct_path, encoding="utf-8") as fh:
            ct = fh.read()
        if 'Extension="xlsx"' not in ct:
            ct = ct.replace(
                "<Default Extension=",
                f'<Default Extension="xlsx" ContentType="{XLSX_CT}"/><Default Extension=',
                1,
            )
            with open(ct_path, "w", encoding="utf-8") as fh:
                fh.write(ct)

        if os.path.exists(dst):
            os.remove(dst)
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in names + [a for a in added if a not in names]:
                zout.write(os.path.join(work, name), name)
        print(f"완료: {dst} (차트 {len(added)}개 변환)")
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    convert(sys.argv[1], sys.argv[2])
