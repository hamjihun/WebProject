"""파일 형식별 텍스트 추출.

각 파서는 Block 목록을 돌려준다. Block 하나는 출처 표시 단위(페이지, 슬라이드,
엑셀 행 묶음 등)이고, 너무 길면 chunker 가 다시 잘게 나눈다.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
import struct
import zipfile
import zlib
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable
from xml.etree import ElementTree as ET


@dataclass
class Block:
    location: str
    text: str
    # 잘게 나눌 때 모든 조각 앞에 다시 붙일 문맥 (예: 엑셀 시트명과 머리글)
    prefix: str = ""


class ParseError(Exception):
    """사용자에게 그대로 보여줄 수 있는 오류."""


SUPPORTED_EXTS = {
    ".pdf", ".docx", ".xlsx", ".xlsm", ".xls", ".csv", ".pptx",
    ".hwp", ".hwpx", ".txt", ".md",
}


def parse_file(path: Path, filename: str, chunk_chars: int = 700) -> list[Block]:
    ext = Path(filename).suffix.lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        if ext in (".ppt", ".doc"):
            raise ParseError(
                f"예전 형식({ext})은 지원하지 않습니다. 오피스에서 "
                f"'{ext}x' 형식으로 다른 이름으로 저장한 뒤 올려주세요."
            )
        raise ParseError(f"지원하지 않는 파일 형식입니다: {ext or '(확장자 없음)'}")
    try:
        blocks = parser(path, chunk_chars)
    except ParseError:
        raise
    except Exception as e:  # 손상된 파일, 암호 걸린 파일 등
        raise ParseError(f"파일을 읽을 수 없습니다 ({type(e).__name__}: {e})") from e
    blocks = [b for b in blocks if len(b.text.strip()) >= 2]
    if not blocks:
        if ext == ".pdf":
            raise ParseError(
                "PDF에서 글자를 찾지 못했습니다. 스캔한 이미지 PDF는 아직 지원하지 않습니다."
            )
        raise ParseError("문서에서 글자를 찾지 못했습니다.")
    return blocks


# ---------------------------------------------------------------- 공통 유틸

def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _fmt_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "예" if v else "아니오"
    if isinstance(v, dt.datetime):
        if v.hour == v.minute == v.second == 0:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e15:
            return f"{int(v):,}" if abs(v) >= 10000 else str(int(v))
        return f"{v:,.4f}".rstrip("0").rstrip(".")
    if isinstance(v, int):
        return f"{v:,}" if abs(v) >= 10000 else str(v)
    return str(v).strip().replace("\n", " ")


# ---------------------------------------------------------------- 표(엑셀/CSV)

def _find_header(rows: list[list[str]]) -> int:
    """앞쪽 10행 중 머리글로 보이는 행의 위치. 없으면 -1."""
    head = rows[:10]
    if not head:
        return -1
    max_filled = max(sum(1 for c in r if c) for r in head)
    if max_filled < 2:
        return -1
    for i, r in enumerate(head):
        filled = [c for c in r if c]
        if len(filled) >= max(2, int(max_filled * 0.6)) and all(
            not re.fullmatch(r"[-+]?[\d,.]+%?", c) for c in filled
        ):
            return i
    return -1


def _col_name(idx: int) -> str:
    name = ""
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        name = chr(65 + rem) + name
    return name


def table_blocks(
    sheet_label: str, rows: list[tuple[int, list[str]]], chunk_chars: int
) -> list[Block]:
    """표를 '머리글: 값' 형태의 행 묶음 Block 들로 바꾼다.

    rows 는 (원본 행 번호, 셀 문자열 목록). 비어 있는 행은 미리 빠져 있어야 한다.
    """
    if not rows:
        return []
    cells = [r for _, r in rows]
    h = _find_header(cells)
    width = max(len(r) for r in cells)
    if h >= 0:
        header = [c or _col_name(i) for i, c in enumerate(cells[h] + [""] * (width - len(cells[h])))]
        title_lines = [" ".join(c for c in r if c) for r in cells[:h]]
        body = rows[h + 1 :]
    else:
        header = [_col_name(i) for i in range(width)]
        title_lines = []
        body = rows

    prefix = f"[{sheet_label}]"
    if title_lines:
        prefix += " " + " / ".join(title_lines)
    prefix += "\n"

    if not body:  # 머리글만 있는 표
        return [Block(sheet_label, prefix + " | ".join(header))]

    blocks: list[Block] = []
    buf: list[str] = []
    first_row = last_row = body[0][0]
    size = 0
    for row_no, r in body:
        pairs = [f"{header[i]}: {c}" for i, c in enumerate(r) if c]
        line = f"({row_no}행) " + " | ".join(pairs)
        if buf and size + len(line) > chunk_chars:
            blocks.append(Block(f"{sheet_label} {first_row}~{last_row}행", "\n".join(buf), prefix))
            buf, size, first_row = [], 0, row_no
        buf.append(line)
        size += len(line) + 1
        last_row = row_no
    if buf:
        loc = f"{sheet_label} {first_row}행" if first_row == last_row else f"{sheet_label} {first_row}~{last_row}행"
        blocks.append(Block(loc, "\n".join(buf), prefix))
    return blocks


def _rows_nonempty(raw_rows: Iterable[Iterable]) -> list[tuple[int, list[str]]]:
    out = []
    for i, row in enumerate(raw_rows, start=1):
        cells = [_fmt_cell(v) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            out.append((i, cells))
    return out


def parse_xlsx(path: Path, chunk_chars: int) -> list[Block]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    blocks: list[Block] = []
    try:
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            rows = _rows_nonempty(ws.iter_rows(values_only=True))
            blocks += table_blocks(f"시트 '{ws.title}'", rows, chunk_chars)
    finally:
        wb.close()
    return blocks


def parse_xls(path: Path, chunk_chars: int) -> list[Block]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    blocks: list[Block] = []
    for sh in book.sheets():
        if sh.visibility != 0:
            continue
        raw = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value if cell.value != "" else None)
            raw.append(row)
        blocks += table_blocks(f"시트 '{sh.name}'", _rows_nonempty(raw), chunk_chars)
    return blocks


def _read_text_file(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp949", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_csv(path: Path, chunk_chars: int) -> list[Block]:
    text = _read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    raw = list(csv.reader(io.StringIO(text), dialect))
    return table_blocks(Path(path).stem, _rows_nonempty(raw), chunk_chars)


# ---------------------------------------------------------------- PPT

def parse_pptx(path: Path, chunk_chars: int) -> list[Block]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))

    def shape_texts(shapes) -> list[str]:
        out: list[str] = []
        # 위→아래, 왼쪽→오른쪽 순서로 읽는다
        ordered = sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
        for sh in ordered:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                out += shape_texts(sh.shapes)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    # 병합된 셀은 한 번만 (빈 칸이 '| |' 로 늘어서는 것 방지)
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells if not c.is_spanned]
                    cells = [c for c in cells if c]
                    if cells:
                        out.append(" | ".join(cells))
            elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    out.append(t)
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                out += _chart_text(sh.chart)
        return out

    blocks: list[Block] = []
    for n, slide in enumerate(prs.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        body = [t for t in shape_texts(slide.shapes) if t != title]
        parts = []
        if title:
            parts.append(f"# {title}")
        parts += body
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append(f"[발표자 노트] {notes}")
        text = _clean("\n".join(parts))
        if text:
            prefix = f"[슬라이드 {n}: {title}]\n" if title else ""
            blocks.append(Block(f"슬라이드 {n}", text, prefix))
    return blocks


def _chart_text(chart) -> list[str]:
    out = []
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            out.append(f"[차트] {chart.chart_title.text_frame.text.strip()}")
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        for s in plot.series:
            vals = [_fmt_cell(v) for v in s.values]
            pairs = ", ".join(f"{c}={v}" for c, v in zip(cats, vals))
            out.append(f"{s.name}: {pairs}")
    except Exception:
        pass
    return out


# ---------------------------------------------------------------- PDF / Word / 텍스트

def parse_pdf(path: Path, chunk_chars: int) -> list[Block]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:
            raise ParseError("암호가 걸린 PDF는 읽을 수 없습니다.") from e
    pages = [_clean(page.extract_text() or "") for page in reader.pages]
    repeated = _repeated_lines(pages)
    blocks = []
    for n, text in enumerate(pages, start=1):
        if repeated:
            text = "\n".join(l for l in text.split("\n") if _norm_line(l) not in repeated).strip()
        if text:
            blocks.append(Block(f"p.{n}", text))
    return blocks


def _norm_line(line: str) -> str:
    line = line.strip()
    # 쪽 번호 줄("- 3 -", "3 / 20", "Page 3 of 5")은 숫자가 달라도 같은 줄로 본다
    if len(re.sub(r"[\d\s\-/|.·]", "", line)) <= 8:
        return re.sub(r"\d+", "#", line)
    return line


def _repeated_lines(pages: list[str]) -> set[str]:
    """여러 페이지의 맨 위·아래에 반복되는 머리글/바닥글 줄."""
    if len(pages) < 4:
        return set()
    counts: dict[str, int] = {}
    for text in pages:
        lines = [l for l in text.split("\n") if l.strip()]
        edge = set(_norm_line(l) for l in lines[:2] + lines[-2:])
        for l in edge:
            counts[l] = counts.get(l, 0) + 1
    return {l for l, c in counts.items() if c >= len(pages) * 0.5 and len(l) < 80}


def parse_docx(path: Path, chunk_chars: int) -> list[Block]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    blocks: list[Block] = []
    heading = ""
    buf: list[str] = []

    def flush():
        if buf:
            loc = f"'{heading[:40]}' 부분" if heading else "본문"
            blocks.append(Block(loc, _clean("\n".join(buf)), f"[{heading}]\n" if heading else ""))
            buf.clear()

    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            if style.lower().startswith(("heading", "title")) or style.startswith(("제목",)):
                flush()
                heading = t
            buf.append(t)
        elif tag == "tbl":
            for row in Table(el, d).rows:
                cells = []
                for c in row.cells:
                    ct = c.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != ct:  # 병합 셀 중복 제거
                        cells.append(ct)
                if any(cells):
                    buf.append(" | ".join(cells))
    flush()
    return blocks


def parse_txt(path: Path, chunk_chars: int) -> list[Block]:
    return [Block("본문", _clean(_read_text_file(path)))]


# ---------------------------------------------------------------- 한글(HWP / HWPX)

def parse_hwpx(path: Path, chunk_chars: int) -> list[Block]:
    with zipfile.ZipFile(path) as z:
        names = sorted(
            (n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)),
            key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[-1]).group(1)),
        )
        parts: list[str] = []
        for name in names:
            with z.open(name) as f:
                line: list[str] = []
                for event, el in ET.iterparse(f, events=("end",)):
                    local = el.tag.rsplit("}", 1)[-1]
                    if local == "t":
                        line.append("".join(el.itertext()))
                    elif local == "p":
                        if line:
                            parts.append("".join(line))
                            line = []
                        el.clear()
                if line:
                    parts.append("".join(line))
    return [Block("본문", _clean("\n".join(parts)))]


HWPTAG_PARA_TEXT = 67
# 8글자(16바이트) 크기를 차지하는 확장/인라인 제어 문자
_HWP_WIDE_CTRLS = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}


def _hwp_para_text(data: bytes) -> str:
    out: list[str] = []
    i, n = 0, len(data) // 2
    while i < n:
        c = struct.unpack_from("<H", data, i * 2)[0]
        if c >= 32:
            out.append(chr(c))
            i += 1
        elif c in _HWP_WIDE_CTRLS:
            if c == 9:
                out.append("\t")
            i += 8
        else:
            if c in (10, 13):
                out.append("\n")
            i += 1
    return "".join(out)


def parse_hwp(path: Path, chunk_chars: int) -> list[Block]:
    import olefile

    if not olefile.isOleFile(str(path)):
        raise ParseError("올바른 한글(HWP) 파일이 아닙니다.")
    ole = olefile.OleFileIO(str(path))
    try:
        header = ole.openstream("FileHeader").read()
        props = struct.unpack_from("<I", header, 36)[0]
        compressed = bool(props & 1)
        if props & 2:
            raise ParseError("암호가 걸린 한글 문서는 읽을 수 없습니다.")
        if props & 4:
            raise ParseError("배포용 한글 문서는 읽을 수 없습니다.")
        sections = sorted(
            (e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText"),
            key=lambda e: int(re.sub(r"\D", "", e[1]) or 0),
        )
        parts: list[str] = []
        for entry in sections:
            data = ole.openstream(entry).read()
            if compressed:
                data = zlib.decompress(data, -15)
            pos = 0
            while pos + 4 <= len(data):
                h = struct.unpack_from("<I", data, pos)[0]
                pos += 4
                tag, size = h & 0x3FF, (h >> 20) & 0xFFF
                if size == 0xFFF:
                    size = struct.unpack_from("<I", data, pos)[0]
                    pos += 4
                if tag == HWPTAG_PARA_TEXT:
                    parts.append(_hwp_para_text(data[pos : pos + size]))
                pos += size
    finally:
        ole.close()
    return [Block("본문", _clean("\n".join(parts)))]


_PARSERS: dict[str, Callable[[Path, int], list[Block]]] = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".xlsm": parse_xlsx,
    ".xls": parse_xls,
    ".csv": parse_csv,
    ".pptx": parse_pptx,
    ".hwp": parse_hwp,
    ".hwpx": parse_hwpx,
    ".txt": parse_txt,
    ".md": parse_txt,
}
