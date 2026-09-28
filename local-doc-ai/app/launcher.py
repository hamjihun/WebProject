"""설치형 exe 의 시작점.

- 이미 실행 중이면 브라우저만 다시 연다
- Ollama 가 꺼져 있으면 켠다
- 서버를 띄우고 브라우저를 연 뒤, 작업 표시줄 알림 영역(트레이)에 아이콘을 둔다
- 트레이의 [종료]를 누르면 프로그램이 끝난다
"""
from __future__ import annotations

import argparse
import logging
import socket
import sys
import threading
import time
import webbrowser

import httpx

from . import config

APP_NAME = "사내 문서 AI"
log = logging.getLogger("docai")


def _setup_logging() -> None:
    config.LOG_DIR.mkdir(parents=True, exist_ok=True)
    handler = logging.FileHandler(config.LOG_DIR / "app.log", encoding="utf-8")
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s"))
    root = logging.getLogger()
    root.addHandler(handler)
    root.setLevel(logging.INFO)
    # 창 없는 exe 에서는 stdout/stderr 가 없으므로 로그 파일로 보낸다
    if sys.stdout is None or sys.stderr is None:
        f = open(config.LOG_DIR / "console.log", "a", encoding="utf-8", buffering=1)
        sys.stdout = sys.stdout or f
        sys.stderr = sys.stderr or f


def _is_our_app(port: int) -> bool:
    try:
        r = httpx.get(f"http://{config.HOST}:{port}/api/ping", timeout=1)
        return r.status_code == 200 and r.json().get("app") == "local-doc-ai"
    except Exception:
        return False


def _port_free(port: int) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        try:
            s.bind((config.HOST, port))
            return True
        except OSError:
            return False


def _pick_port() -> tuple[int, bool]:
    """(포트, 이미 실행 중인지)"""
    for port in range(config.PORT, config.PORT + 10):
        if _is_our_app(port):
            return port, True
        if _port_free(port):
            return port, False
    raise RuntimeError("사용할 수 있는 포트가 없습니다.")


def _ensure_ollama() -> None:
    from . import ollama_client, system

    try:
        ollama_client.list_models()
        return
    except ollama_client.OllamaError:
        pass
    if system.start_ollama():
        log.info("Ollama 시작")


def _message_box(text: str) -> None:
    if sys.platform == "win32":
        import ctypes

        ctypes.windll.user32.MessageBoxW(None, text, APP_NAME, 0x40)
    else:
        print(text)


def _run_tray(url: str, stop: threading.Event) -> None:
    try:
        import pystray
        from PIL import Image, ImageDraw
    except ImportError:
        stop.wait()
        return

    img = Image.new("RGBA", (64, 64), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((2, 2, 62, 62), 14, fill=(47, 91, 211))
    d.polygon([(18, 14), (38, 14), (46, 22), (46, 50), (18, 50)], fill="white")
    d.polygon([(38, 14), (38, 22), (46, 22)], fill=(200, 212, 245))

    def on_open(icon, item):
        webbrowser.open(url)

    def on_quit(icon, item):
        stop.set()
        icon.stop()

    icon = pystray.Icon(
        "local-doc-ai", img, APP_NAME,
        menu=pystray.Menu(
            pystray.MenuItem("열기", on_open, default=True),
            pystray.MenuItem("종료", on_quit),
        ),
    )
    threading.Thread(target=lambda: (stop.wait(), icon.stop()), daemon=True).start()
    icon.run()


def _selftest_parsers() -> None:
    """exe 안에 문서 라이브러리의 템플릿 파일이 빠짐없이 들어갔는지 확인."""
    import tempfile
    from pathlib import Path

    import docx
    import openpyxl
    from pptx import Presentation

    from .parsers import parse_file

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "시험 슬라이드"
        prs.save(tmp / "a.pptx")
        d = docx.Document()
        d.add_paragraph("시험 문단")
        d.save(tmp / "a.docx")
        wb = openpyxl.Workbook()
        wb.active.append(["항목", "값"])
        wb.active.append(["시험", 1])
        wb.save(tmp / "a.xlsx")
        for name in ("a.pptx", "a.docx", "a.xlsx"):
            assert parse_file(tmp / name, name), name


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--selftest", action="store_true", help="서버를 띄워 확인만 하고 종료 (빌드 검사용)")
    parser.add_argument("--no-browser", action="store_true")
    args = parser.parse_args(argv)

    _setup_logging()
    port, running = _pick_port()
    url = f"http://{config.HOST}:{port}"
    if running and not args.selftest:
        webbrowser.open(url)
        return 0

    import uvicorn

    from .main import app

    threading.Thread(target=_ensure_ollama, daemon=True).start()

    server = uvicorn.Server(uvicorn.Config(app, host=config.HOST, port=port, log_config=None, log_level="warning"))
    t = threading.Thread(target=server.run, daemon=True)
    t.start()
    for _ in range(100):
        if server.started:
            break
        time.sleep(0.1)
    if not server.started:
        _message_box("프로그램을 시작하지 못했습니다. 로그 폴더를 확인해 주세요:\n" + str(config.LOG_DIR))
        return 1
    log.info("started on %s", url)

    if args.selftest:
        ok = True
        report = []
        for path in ("/api/ping", "/", "/static/app.js", "/api/notebooks", "/api/setup"):
            r = httpx.get(url + path, timeout=10)
            report.append(f"{path} {r.status_code}")
            ok &= r.status_code == 200
        try:
            _selftest_parsers()
            report.append("parsers ok")
        except Exception as e:
            report.append(f"parsers FAIL {e!r}")
            ok = False
        server.should_exit = True
        t.join(10)
        report.append("SELFTEST " + ("OK" if ok else "FAIL"))
        # 창 없는 exe 는 화면 출력이 없으므로 파일로도 남긴다
        (config.LOG_DIR / "selftest.txt").write_text("\n".join(report), encoding="utf-8")
        print("\n".join(report))
        return 0 if ok else 1

    if not args.no_browser:
        webbrowser.open(url)

    stop = threading.Event()
    try:
        _run_tray(url, stop)
    except Exception:
        log.exception("tray failed")
        stop.wait()
    server.should_exit = True
    t.join(10)
    return 0


if __name__ == "__main__":
    sys.exit(main())
