import datetime as dt
import struct
import zipfile

import pytest

from app.chunker import chunk_blocks
from app.parsers import Block, ParseError, _hwp_para_text, parse_file


def test_xlsx_title_rows_header_and_row_numbers(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2분기"
    ws.append(["2024년 2분기 매출 현황"])
    ws.append([])
    ws.append(["거래처", "품목", "금액", "일자"])
    ws.append(["한빛상사", "볼트", 1250000, dt.datetime(2024, 4, 3)])
    ws.append(["누리산업", "너트", 30.5, dt.datetime(2024, 5, 1)])
    hidden = wb.create_sheet("숨김")
    hidden.append(["비밀", "값"])
    hidden.sheet_state = "hidden"
    p = tmp_path / "매출.xlsx"
    wb.save(p)

    blocks = parse_file(p, "매출.xlsx")
    assert len(blocks) == 1
    b = blocks[0]
    assert b.location == "시트 '2분기' 4~5행"
    assert "2024년 2분기 매출 현황" in b.prefix
    assert "(4행) 거래처: 한빛상사 | 품목: 볼트 | 금액: 1,250,000 | 일자: 2024-04-03" in b.text
    assert "금액: 30.5" in b.text
    assert "비밀" not in b.text + b.prefix


def test_xlsx_large_sheet_is_grouped_and_prefix_repeated(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["사번", "이름", "부서"])
    for i in range(200):
        ws.append([1000 + i, f"직원{i}", "영업팀"])
    p = tmp_path / "인사.xlsx"
    wb.save(p)

    blocks = parse_file(p, "인사.xlsx", chunk_chars=500)
    assert len(blocks) > 5
    assert all(len(b.text) <= 500 for b in blocks)
    chunks = chunk_blocks(blocks, 500)
    assert all(c.text.startswith("[시트 'Sheet']") for c in chunks)
    assert blocks[0].location.startswith("시트 'Sheet' 2~")


def test_csv_cp949(tmp_path):
    p = tmp_path / "목록.csv"
    p.write_bytes("이름,전화\n홍길동,010-1234\n".encode("cp949"))
    blocks = parse_file(p, "목록.csv")
    assert "이름: 홍길동 | 전화: 010-1234" in blocks[0].text


def test_xls(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("재고")
    for c, v in enumerate(["품번", "수량"]):
        ws.write(0, c, v)
    ws.write(1, 0, "A-100")
    ws.write(1, 1, 42)
    p = tmp_path / "재고.xls"
    wb.save(str(p))
    blocks = parse_file(p, "재고.xls")
    assert "품번: A-100 | 수량: 42" in blocks[0].text


def test_pptx_slides_tables_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "영업 전략"
    s1.placeholders[1].text = "신규 거래처 20곳 확보"
    s1.notes_slide.notes_text_frame.text = "3분기까지 완료"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "예산"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "항목", "금액"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "광고", "500만원"
    p = tmp_path / "전략.pptx"
    prs.save(p)

    blocks = parse_file(p, "전략.pptx")
    assert [b.location for b in blocks] == ["슬라이드 1", "슬라이드 2"]
    assert "# 영업 전략" in blocks[0].text
    assert "신규 거래처 20곳 확보" in blocks[0].text
    assert "[발표자 노트] 3분기까지 완료" in blocks[0].text
    assert "광고 | 500만원" in blocks[1].text


def test_docx_headings_and_tables(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("휴가 규정", 1)
    d.add_paragraph("연차는 15일이다.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "구분", "일수"
    p = tmp_path / "규정.docx"
    d.save(p)
    blocks = parse_file(p, "규정.docx")
    assert blocks[0].location == "'휴가 규정' 부분"
    assert "연차는 15일이다." in blocks[0].text
    assert "구분 | 일수" in blocks[0].text


def _make_pdf(pages: list[str]) -> bytes:
    """글꼴 하나로 된 간단한 텍스트 PDF를 직접 만든다."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        lines = text.split("\n")
        ops = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({l}) Tj T*" for l in lines) + " ET"
        objs.append(f"<< /Length {len(ops)} >>\nstream\n{ops}\nendstream")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {len(objs)} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(len(objs))
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_pdf_pages_and_repeated_footer_removed(tmp_path):
    pages = [f"ACME Corp Confidential\nSection {i} budget is {i * 100} dollars\nPage {i} of 5" for i in range(1, 6)]
    p = tmp_path / "report.pdf"
    p.write_bytes(_make_pdf(pages))
    blocks = parse_file(p, "report.pdf")
    assert [b.location for b in blocks] == [f"p.{i}" for i in range(1, 6)]
    assert "Section 3 budget is 300 dollars" in blocks[2].text
    assert "Confidential" not in blocks[2].text
    assert "of 5" not in blocks[2].text


def test_pdf_without_text_gives_friendly_error(tmp_path):
    p = tmp_path / "scan.pdf"
    p.write_bytes(_make_pdf([""]))
    with pytest.raises(ParseError, match="스캔"):
        parse_file(p, "scan.pdf")


def test_hwpx(tmp_path):
    xml = ('<?xml version="1.0" encoding="UTF-8"?>'
           '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
           'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
           '<hp:p><hp:run><hp:t>출장비 규정</hp:t></hp:run></hp:p>'
           '<hp:p><hp:run><hp:t>일비는 </hp:t><hp:t>5만원</hp:t></hp:run></hp:p></hs:sec>')
    p = tmp_path / "규정.hwpx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("Contents/section0.xml", xml)
    blocks = parse_file(p, "규정.hwpx")
    assert blocks[0].text == "출장비 규정\n일비는 5만원"


def test_hwp_para_text_skips_controls():
    # "가" + 확장 제어문자(8글자 크기) + "나" + 줄바꿈 + "다"
    chars = [ord("가"), 11, 0, 0, 0, 0, 0, 0, 11, ord("나"), 13, ord("다")]
    data = struct.pack(f"<{len(chars)}H", *chars)
    assert _hwp_para_text(data) == "가나\n다"


def test_unsupported_and_legacy_formats(tmp_path):
    p = tmp_path / "x.ppt"
    p.write_bytes(b"x")
    with pytest.raises(ParseError, match="pptx"):
        parse_file(p, "x.ppt")
    with pytest.raises(ParseError, match="지원하지 않는"):
        parse_file(p, "x.zip")


def test_chunker_splits_long_text_with_overlap():
    text = "\n".join(f"{i}번째 문장은 규정에 관한 설명입니다." for i in range(100))
    chunks = chunk_blocks([Block("p.1", text, "[머리]\n")], chunk_chars=300, overlap=60)
    assert len(chunks) > 5
    assert all(len(c.text) <= 300 for c in chunks)
    assert all(c.text.startswith("[머리]\n") for c in chunks)
    first_last = chunks[0].text.split("\n")[-1]
    assert first_last in chunks[1].text  # 겹침
